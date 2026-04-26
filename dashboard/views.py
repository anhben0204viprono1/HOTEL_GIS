"""
dashboard/views.py
Trang admin tự xây — yêu cầu staff / superuser / group Employee.
"""
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.db.models import Count, Sum, Avg, Q
from django.db import transaction
from django.utils import timezone
from django.http import JsonResponse, HttpResponse
from django.utils.text import slugify
from django.db.models.functions import TruncMonth
from datetime import timedelta, date
import json
from decimal import Decimal, InvalidOperation
import datetime

from hotels.models import Hotel, RoomType, Room, Amenity, HotelImage, RoomTypeImage, HomepageConfig, HotelService
from bookings.models import Booking, Payment, Review
from django.contrib.auth.models import User, Group
from hotels.widgets import LeafletMapWidget
from staff.models import StaffProfile


# ── Guard: staff / superuser / group Employee mới vào được ──────────────────
def is_staff(user):
    return user.is_authenticated and (
        user.is_staff or user.is_superuser or
        user.groups.filter(name='Employee').exists()
    )

staff_required = user_passes_test(is_staff, login_url='/accounts/login/')
superuser_required = user_passes_test(
    lambda u: u.is_superuser, login_url='/accounts/login/'
)


# ═══════════════════════════════════════════════════════════
# DASHBOARD — Trang chủ admin
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def dashboard_home(request):
    today = date.today()
    month_start = today.replace(day=1)

    stats = {
        'total_hotels':   Hotel.objects.filter(is_active=True).count(),
        'total_rooms':    Room.objects.filter(status='available').count(),
        'total_users':    User.objects.filter(is_active=True, is_staff=False).count(),
        'total_bookings': Booking.objects.count(),

        'bookings_month': Booking.objects.filter(created_at__gte=month_start).count(),
        'revenue_month':  Payment.objects.filter(
            status='paid', paid_at__gte=month_start
        ).aggregate(total=Sum('amount'))['total'] or 0,

        'checkin_today':  Booking.objects.filter(check_in=today, status='confirmed').count(),
        'checkout_today': Booking.objects.filter(check_out=today, status='checked_in').count(),
        'pending':        Booking.objects.filter(status='pending').count(),
    }

    chart_labels, chart_data = [], []
    for i in range(6, -1, -1):
        d = today - timedelta(days=i)
        chart_labels.append(d.strftime('%d/%m'))
        chart_data.append(Booking.objects.filter(created_at__date=d).count())

    recent_bookings = (
        Booking.objects
        .select_related('user', 'room__room_type__hotel')
        .order_by('-created_at')[:8]
    )
    top_hotels = (
        Hotel.objects
        .annotate(booking_count=Count('rooms__bookings'))
        .order_by('-booking_count')[:5]
    )

    return render(request, 'dashboard/home.html', {
        'stats':           stats,
        'chart_labels':    json.dumps(chart_labels),
        'chart_data':      json.dumps(chart_data),
        'recent_bookings': recent_bookings,
        'top_hotels':      top_hotels,
        'page': 'dashboard',
    })


# ═══════════════════════════════════════════════════════════
# HOTELS — Quản lý khách sạn
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def hotel_list(request):
    q      = request.GET.get('q', '')
    city   = request.GET.get('city', '')
    hotels = Hotel.objects.annotate(
        room_count=Count('rooms', distinct=True),
        booking_count=Count('rooms__bookings', distinct=True),
    ).order_by('-created_at')

    if q:
        hotels = hotels.filter(Q(name__icontains=q) | Q(city__icontains=q))
    if city:
        hotels = hotels.filter(city__icontains=city)

    cities = Hotel.objects.values_list('city', flat=True).distinct().order_by('city')
    return render(request, 'dashboard/hotels/list.html', {
        'hotels': hotels,
        'filters': {'q': q, 'city': city},
        'cities': cities, 'page': 'hotels',
    })


def _openpyxl_or_message(request):
    try:
        import openpyxl  # type: ignore
    except Exception:
        openpyxl = None
    if openpyxl is None:
        messages.error(
            request,
            'Thiếu thư viện openpyxl. Hãy chạy: pip install -r requirements.txt',
        )
    return openpyxl


def _normalize_header(value: str) -> str:
    return (
        value.strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )


def _as_decimal(value):
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    s = str(value).strip().replace(",", ".")
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _as_time(value, default_value):
    if value is None or value == "":
        return default_value
    if isinstance(value, datetime.time):
        return value
    if isinstance(value, datetime.datetime):
        return value.time()
    s = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    return default_value


@login_required
@superuser_required
def hotel_import_template(request):
    openpyxl = _openpyxl_or_message(request)
    if openpyxl is None:
        return redirect('dashboard:hotel_list')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hotels"
    headers = [
        "id(optional)",
        "name*",
        "city",
        "address",
        "latitude*",
        "longitude*",
        "star_rating(1-5)",
        "phone",
        "email",
        "short_description",
        "description",
        "thumbnail_url",
        "website",
        "check_in_time(HH:MM)",
        "check_out_time(HH:MM)",
        "is_active(TRUE/FALSE)",
    ]
    ws.append(headers)
    ws.append([
        "",
        "Hotel Demo",
        "Hồ Chí Minh",
        "123 Nguyễn Huệ, Q1",
        "10.7769",
        "106.7009",
        4,
        "0900000000",
        "demo@example.com",
        "Mô tả ngắn",
        "Mô tả dài",
        "",
        "https://example.com",
        "14:00",
        "12:00",
        True,
    ])

    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = 'attachment; filename="hotel_import_template.xlsx"'
    return resp


@login_required
@superuser_required
def hotel_import_excel(request):
    if request.method == 'GET':
        return render(request, 'dashboard/hotels/import.html', {'page': 'hotels'})

    openpyxl = _openpyxl_or_message(request)
    if openpyxl is None:
        return redirect('dashboard:hotel_list')

    f = request.FILES.get('file')
    if not f:
        messages.error(request, 'Vui lòng chọn file Excel (.xlsx).')
        return redirect('dashboard:hotel_import_excel')

    try:
        wb = openpyxl.load_workbook(f, data_only=True)
    except Exception:
        messages.error(request, 'Không đọc được file. Hãy chắc chắn file là .xlsx hợp lệ.')
        return redirect('dashboard:hotel_import_excel')

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        messages.error(request, 'File Excel rỗng.')
        return redirect('dashboard:hotel_import_excel')

    raw_headers = [str(c or "").strip() for c in rows[0]]
    normalized = [_normalize_header(h) for h in raw_headers]

    aliases = {
        "id(optional)": "id",
        "id": "id",
        "pk": "id",
        "name*": "name",
        "tenkhachsan": "name",
        "ten": "name",
        "name": "name",
        "city": "city",
        "thanhpho": "city",
        "address": "address",
        "diachi": "address",
        "latitude*": "latitude",
        "latitude": "latitude",
        "vido": "latitude",
        "longitude*": "longitude",
        "longitude": "longitude",
        "kinhdo": "longitude",
        "starrating(1-5)": "star_rating",
        "starrating": "star_rating",
        "sosao": "star_rating",
        "phone": "phone",
        "sodienthoai": "phone",
        "email": "email",
        "shortdescription": "short_description",
        "motangan": "short_description",
        "description": "description",
        "motadai": "description",
        "thumbnailurl": "thumbnail_url",
        "website": "website",
        "checkintime(hh:mm)": "check_in_time",
        "checkintime": "check_in_time",
        "checkouttime(hh:mm)": "check_out_time",
        "checkouttime": "check_out_time",
        "isactive(true/false)": "is_active",
        "isactive": "is_active",
    }

    field_to_idx = {}
    for idx, h in enumerate(normalized):
        key = aliases.get(h)
        if key:
            field_to_idx[key] = idx

    missing = [k for k in ("name", "latitude", "longitude") if k not in field_to_idx]
    if missing:
        messages.error(
            request,
            f"Thiếu cột bắt buộc: {', '.join(missing)}. Tải template để đúng format.",
        )
        return redirect('dashboard:hotel_import_excel')

    created = 0
    updated = 0
    errors = []

    with transaction.atomic():
        for r_i, row in enumerate(rows[1:], start=2):
            name = row[field_to_idx["name"]]
            if name is None or str(name).strip() == "":
                continue

            lat = _as_decimal(row[field_to_idx["latitude"]])
            lng = _as_decimal(row[field_to_idx["longitude"]])
            if lat is None or lng is None:
                errors.append(f"Dòng {r_i}: latitude/longitude không hợp lệ.")
                continue

            hotel_id = None
            if "id" in field_to_idx:
                v = row[field_to_idx["id"]]
                try:
                    hotel_id = int(v) if v not in (None, "") else None
                except Exception:
                    hotel_id = None

            if hotel_id:
                hotel = Hotel.objects.filter(pk=hotel_id).first()
            else:
                hotel = None

            is_new = hotel is None
            if hotel is None:
                hotel = Hotel()

            hotel.name = str(name).strip()
            if "city" in field_to_idx and row[field_to_idx["city"]] not in (None, ""):
                hotel.city = str(row[field_to_idx["city"]]).strip()
            if "address" in field_to_idx and row[field_to_idx["address"]] not in (None, ""):
                hotel.address = str(row[field_to_idx["address"]]).strip()

            hotel.latitude = lat
            hotel.longitude = lng

            if "star_rating" in field_to_idx and row[field_to_idx["star_rating"]] not in (None, ""):
                try:
                    sr = int(row[field_to_idx["star_rating"]])
                    hotel.star_rating = min(max(sr, 1), 5)
                except Exception:
                    pass
            if "phone" in field_to_idx and row[field_to_idx["phone"]] not in (None, ""):
                hotel.phone = str(row[field_to_idx["phone"]]).strip()
            if "email" in field_to_idx and row[field_to_idx["email"]] not in (None, ""):
                hotel.email = str(row[field_to_idx["email"]]).strip()
            if "short_description" in field_to_idx and row[field_to_idx["short_description"]] not in (None, ""):
                hotel.short_description = str(row[field_to_idx["short_description"]]).strip()
            if "description" in field_to_idx and row[field_to_idx["description"]] not in (None, ""):
                hotel.description = str(row[field_to_idx["description"]]).strip()
            if "thumbnail_url" in field_to_idx and row[field_to_idx["thumbnail_url"]] not in (None, ""):
                hotel.thumbnail_url = str(row[field_to_idx["thumbnail_url"]]).strip()
            if "website" in field_to_idx and row[field_to_idx["website"]] not in (None, ""):
                hotel.website = str(row[field_to_idx["website"]]).strip()

            hotel.check_in_time = _as_time(
                row[field_to_idx["check_in_time"]] if "check_in_time" in field_to_idx else None,
                hotel.check_in_time,
            )
            hotel.check_out_time = _as_time(
                row[field_to_idx["check_out_time"]] if "check_out_time" in field_to_idx else None,
                hotel.check_out_time,
            )

            if "is_active" in field_to_idx:
                v = row[field_to_idx["is_active"]]
                if isinstance(v, bool):
                    hotel.is_active = v
                elif v is not None and str(v).strip() != "":
                    s = str(v).strip().lower()
                    hotel.is_active = s in {"1", "true", "yes", "y", "on"}

            if is_new:
                base = slugify(hotel.name, allow_unicode=True)[:250] or "hotel"
                slug = base
                i = 2
                while Hotel.objects.filter(slug=slug).exists():
                    slug = f"{base}-{i}"
                    i += 1
                hotel.slug = slug

            try:
                hotel.save()
            except Exception as e:
                errors.append(f"Dòng {r_i}: không lưu được ({e}).")
                continue

            if is_new:
                created += 1
            else:
                updated += 1

    if created or updated:
        messages.success(request, f"✅ Import xong: tạo mới {created}, cập nhật {updated}.")
    if errors:
        messages.warning(request, f"⚠ Có {len(errors)} dòng lỗi (xem chi tiết bên dưới).")
    return render(
        request,
        'dashboard/hotels/import.html',
        {'created': created, 'updated': updated, 'errors': errors, 'page': 'hotels'},
    )


@login_required
@staff_required
def revenue_export_excel(request):
    openpyxl = _openpyxl_or_message(request)
    if openpyxl is None:
        return redirect('dashboard:home')

    try:
        year = int(request.GET.get('year') or timezone.localdate().year)
    except Exception:
        year = timezone.localdate().year

    qs = (
        Payment.objects
        .filter(status='paid', paid_at__isnull=False, paid_at__year=year)
        .annotate(month=TruncMonth('paid_at'))
        .values('month')
        .annotate(
            total=Sum('amount'),
            payment_count=Count('id'),
            booking_count=Count('booking', distinct=True),
        )
        .order_by('month')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Revenue {year}"
    ws.append(["Tháng", "Doanh thu (VNĐ)", "Số giao dịch", "Số booking"])

    total_year = Decimal("0")
    for row in qs:
        month = row["month"]
        total = row["total"] or 0
        total_year += Decimal(str(total))
        ws.append([
            month.strftime("%m/%Y") if month else "",
            float(total),
            int(row["payment_count"] or 0),
            int(row["booking_count"] or 0),
        ])

    ws.append([])
    ws.append(["Tổng năm", float(total_year), "", ""])

    for cell in ws["B"][1:]:
        cell.number_format = '#,##0'

    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="revenue_{year}.xlsx"'
    return resp


@login_required
@staff_required
def hotel_create(request):
    hotel = Hotel()
    return _save_hotel(request, hotel)


@login_required
@staff_required
def hotel_edit(request, pk):
    hotel = get_object_or_404(Hotel, pk=pk)
    return _save_hotel(request, hotel)


@login_required
@staff_required
def hotel_delete(request, pk):
    hotel = get_object_or_404(Hotel, pk=pk)
    if request.method == 'POST':
        hotel.delete()
        messages.success(request, f'Đã xóa khách sạn "{hotel.name}".')
    return redirect('dashboard:hotel_list')


@login_required
@staff_required
def hotel_toggle(request, pk):
    hotel = get_object_or_404(Hotel, pk=pk)
    hotel.is_active = not hotel.is_active
    hotel.save()
    status = 'kích hoạt' if hotel.is_active else 'tắt'
    messages.success(request, f'Đã {status} khách sạn "{hotel.name}".')
    return redirect('dashboard:hotel_list')


def _save_hotel(request, hotel):
    amenities_all = Amenity.objects.all().order_by('category', 'name')
    map_widget    = LeafletMapWidget()
    is_edit       = bool(hotel.pk)

    if request.method == 'POST':
        hotel.name         = request.POST.get('name', '').strip()
        hotel.slug         = request.POST.get('slug', '').strip() or None
        hotel.address      = request.POST.get('address', '').strip()
        hotel.city         = request.POST.get('city', '').strip() or 'Hồ Chí Minh'
        hotel.phone        = request.POST.get('phone', '').strip()
        hotel.email        = request.POST.get('email', '').strip()
        hotel.star_rating  = int(request.POST.get('star_rating', 3))
        hotel.short_description = request.POST.get('short_description', '').strip()
        hotel.description  = request.POST.get('description', '').strip()
        hotel.thumbnail_url= request.POST.get('thumbnail_url', '').strip()
        hotel.website      = request.POST.get('website', '').strip()
        hotel.check_in_time  = request.POST.get('check_in_time', '14:00') or '14:00'
        hotel.check_out_time = request.POST.get('check_out_time', '12:00') or '12:00'
        hotel.is_active    = 'is_active' in request.POST

        lat = request.POST.get('latitude', '').strip()
        lng = request.POST.get('longitude', '').strip()
        try:
            hotel.latitude  = float(lat)
            hotel.longitude = float(lng)
        except (ValueError, TypeError):
            messages.error(request, 'Tọa độ không hợp lệ.')
            return render(request, 'dashboard/hotels/form.html', {
                'hotel': hotel, 'amenities_all': amenities_all,
                'map_widget': map_widget, 'is_edit': is_edit, 'page': 'hotels',
            })

        if 'image' in request.FILES:
            hotel.image = request.FILES['image']

        try:
            hotel.save()
        except Exception as e:
            messages.error(request, f'Lỗi lưu khách sạn: {e}')
            return render(request, 'dashboard/hotels/form.html', {
                'hotel': hotel, 'amenities_all': amenities_all,
                'map_widget': map_widget, 'is_edit': is_edit, 'page': 'hotels',
            })

        # Tiện nghi cấp khách sạn
        selected_amenity_ids = request.POST.getlist('hotel_amenities')
        hotel.amenities.set(Amenity.objects.filter(id__in=selected_amenity_ids))

        _save_hotel_gallery_images(request, hotel)

        action = 'Cập nhật' if is_edit else 'Thêm mới'
        messages.success(request, f'{action} khách sạn "{hotel.name}" thành công.')
        return redirect('dashboard:hotel_list')

    return render(request, 'dashboard/hotels/form.html', {
        'hotel': hotel, 'amenities_all': amenities_all,
        'map_widget': map_widget, 'is_edit': is_edit, 'page': 'hotels',
        'hotel_amenity_ids': list(hotel.amenities.values_list('id', flat=True)) if hotel.pk else [],
    })


def _save_hotel_gallery_images(request, hotel):
    """Xử lý upload nhiều ảnh gallery."""
    gallery_files = request.FILES.getlist('gallery_images')
    captions      = request.POST.getlist('gallery_captions')
    delete_ids    = request.POST.getlist('delete_image_ids')

    if delete_ids:
        HotelImage.objects.filter(hotel=hotel, id__in=delete_ids).delete()

    for i, f in enumerate(gallery_files):
        caption = captions[i] if i < len(captions) else ''
        HotelImage.objects.create(hotel=hotel, image=f, caption=caption)


# ═══════════════════════════════════════════════════════════
# BOOKINGS — Quản lý đặt phòng
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def booking_list(request):
    status = request.GET.get('status', '')
    q      = request.GET.get('q', '')
    hotel  = request.GET.get('hotel', '')

    bookings = (
        Booking.objects
        .select_related('user', 'room__room_type__hotel')
        .order_by('-created_at')
    )
    if status:
        bookings = bookings.filter(status=status)
    if hotel:
        bookings = bookings.filter(room__room_type__hotel_id=hotel)
    if q:
        bookings = bookings.filter(
            Q(user__username__icontains=q) |
            Q(user__email__icontains=q) |
            Q(room__room_number__icontains=q)
        )

    hotels_qs = Hotel.objects.filter(is_active=True).order_by('name')
    return render(request, 'dashboard/bookings/list.html', {
        'bookings': bookings, 'status': status, 'q': q,
        'hotel': hotel, 'hotels': hotels_qs,
        'status_choices': Booking.STATUS_CHOICES,
        'page': 'bookings',
    })


@login_required
@staff_required
def booking_detail(request, pk):
    booking = get_object_or_404(
        Booking.objects.select_related('user', 'room__room_type__hotel'),
        pk=pk
    )
    amenity_usages = booking.amenity_usages.select_related('amenity').all()
    return render(request, 'dashboard/bookings/detail.html', {
        'booking': booking,
        'amenity_usages': amenity_usages,
        'page': 'bookings',
    })


@login_required
@staff_required
def booking_update_status(request, pk):
    booking = get_object_or_404(Booking, pk=pk)
    if request.method == 'POST':
        new_status = request.POST.get('status')
        valid = [s[0] for s in Booking.STATUS_CHOICES]
        if new_status in valid:
            booking.status = new_status
            booking.save()
            messages.success(request, f'Đã cập nhật trạng thái booking #{pk}.')
        else:
            messages.error(request, 'Trạng thái không hợp lệ.')
    return redirect('dashboard:booking_detail', pk=pk)


# ═══════════════════════════════════════════════════════════
# ROOMS — Quản lý phòng
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def room_list(request):
    status = request.GET.get('status', '')
    hotel  = request.GET.get('hotel', '')
    rooms  = Room.objects.select_related('room_type__hotel').order_by('room_type__hotel', 'room_number')

    if status:
        rooms = rooms.filter(status=status)
    if hotel:
        rooms = rooms.filter(room_type__hotel_id=hotel)

    hotels_qs = Hotel.objects.filter(is_active=True).order_by('name')
    return render(request, 'dashboard/rooms/list.html', {
        'rooms': rooms, 'status': status, 'hotel': hotel,
        'hotels': hotels_qs, 'status_choices': Room.STATUS_CHOICES, 'page': 'rooms',
    })


@login_required
@staff_required
def room_update_status(request, pk):
    room = get_object_or_404(Room, pk=pk)
    if request.method == 'POST':
        new_status = request.POST.get('status')
        valid = [s[0] for s in Room.STATUS_CHOICES]
        if new_status in valid:
            room.status = new_status
            room.save()
            messages.success(request, f'Đã cập nhật trạng thái phòng {room.room_number}.')
    return redirect('dashboard:room_list')


# ═══════════════════════════════════════════════════════════
# USERS — Quản lý tài khoản
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def user_list(request):
    q = request.GET.get('q', '')
    users = User.objects.annotate(
        booking_count=Count('bookings')
    ).order_by('-date_joined')

    if q:
        users = users.filter(
            Q(username__icontains=q) |
            Q(email__icontains=q) |
            Q(first_name__icontains=q)
        )
    return render(request, 'dashboard/users/list.html', {
        'users': users, 'q': q, 'page': 'users',
    })


@login_required
@staff_required
def user_toggle(request, pk):
    user = get_object_or_404(User, pk=pk)
    if user != request.user:
        user.is_active = not user.is_active
        user.save()
        status = 'kích hoạt' if user.is_active else 'khóa'
        messages.success(request, f'Đã {status} tài khoản "{user.username}".')
    return redirect('dashboard:user_list')


@login_required
@superuser_required
def user_permissions(request, pk):
    """Chỉ superuser mới chỉnh được quyền user."""
    target_user = get_object_or_404(User, pk=pk)
    groups_all  = Group.objects.all()

    if request.method == 'POST':
        selected_ids = request.POST.getlist('groups')
        target_user.groups.set(Group.objects.filter(id__in=selected_ids))
        target_user.is_staff     = 'is_staff'     in request.POST
        target_user.is_superuser = 'is_superuser' in request.POST
        target_user.save()
        messages.success(request, f'Đã cập nhật quyền cho "{target_user.username}".')
        return redirect('dashboard:user_list')

    return render(request, 'dashboard/users/permissions.html', {
        'target_user':     target_user,
        'groups_all':      groups_all,
        'user_group_ids':  list(target_user.groups.values_list('id', flat=True)),
        'page': 'users',
    })


# ═══════════════════════════════════════════════════════════
# ROOM TYPES
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def roomtype_list(request, hotel_pk):
    hotel = get_object_or_404(Hotel, pk=hotel_pk)
    room_types = hotel.room_types.annotate(
        room_count=Count('rooms')
    ).order_by('price_per_night')
    return render(request, 'dashboard/rooms/roomtype_list.html', {
        'hotel': hotel, 'room_types': room_types, 'page': 'hotels',
    })


@login_required
@staff_required
def roomtype_create(request, hotel_pk):
    hotel     = get_object_or_404(Hotel, pk=hotel_pk)
    room_type = RoomType(hotel=hotel)
    return _save_roomtype(request, hotel, room_type)


@login_required
@staff_required
def roomtype_edit(request, hotel_pk, pk):
    hotel     = get_object_or_404(Hotel, pk=hotel_pk)
    room_type = get_object_or_404(RoomType, pk=pk, hotel=hotel)
    return _save_roomtype(request, hotel, room_type)


@login_required
@staff_required
def roomtype_delete(request, hotel_pk, pk):
    hotel     = get_object_or_404(Hotel, pk=hotel_pk)
    room_type = get_object_or_404(RoomType, pk=pk, hotel=hotel)
    if request.method == 'POST':
        room_type.delete()
        messages.success(request, f'Đã xóa loại phòng "{room_type.name}".')
    return redirect('dashboard:roomtype_list', hotel_pk=hotel_pk)


def _save_roomtype(request, hotel, room_type):
    amenities_all = Amenity.objects.all().order_by('category', 'name')
    is_edit = bool(room_type.pk)

    if request.method == 'POST':
        room_type.name          = request.POST.get('name', '').strip()
        room_type.description   = request.POST.get('description', '').strip()
        room_type.bed_type      = request.POST.get('bed_type', '')
        room_type.is_active     = 'is_active' in request.POST

        try:
            room_type.max_occupancy   = int(request.POST.get('max_occupancy', 2))
            room_type.price_per_night = float(request.POST.get('price_per_night', 0))
            # Giá theo giờ (có thể để trống)
            pph = request.POST.get('price_per_hour', '').strip()
            room_type.price_per_hour = float(pph) if pph else None
            area = request.POST.get('area_sqm', '').strip()
            room_type.area_sqm = float(area) if area else None
        except (ValueError, TypeError) as e:
            messages.error(request, f'Dữ liệu số không hợp lệ: {e}')
            return render(request, 'dashboard/rooms/roomtype_form.html', {
                'hotel': hotel, 'room_type': room_type,
                'amenities_all': amenities_all, 'is_edit': is_edit, 'page': 'hotels',
            })

        room_type.thumbnail_url = request.POST.get('thumbnail_url', '').strip()
        if 'image' in request.FILES:
            room_type.image = request.FILES['image']

        try:
            room_type.save()
        except Exception as e:
            messages.error(request, f'Lỗi lưu loại phòng: {e}')
            return render(request, 'dashboard/rooms/roomtype_form.html', {
                'hotel': hotel, 'room_type': room_type,
                'amenities_all': amenities_all, 'is_edit': is_edit, 'page': 'hotels',
            })

        # Amenities
        selected_ids = request.POST.getlist('amenities')
        room_type.amenities.set(Amenity.objects.filter(id__in=selected_ids))

        # Gallery ảnh loại phòng
        gallery_files = request.FILES.getlist('gallery_images')
        delete_ids    = request.POST.getlist('delete_image_ids')
        if delete_ids:
            RoomTypeImage.objects.filter(room_type=room_type, id__in=delete_ids).delete()
        for f in gallery_files:
            RoomTypeImage.objects.create(room_type=room_type, image=f)

        action = 'Cập nhật' if is_edit else 'Thêm mới'
        messages.success(request, f'{action} loại phòng "{room_type.name}" thành công.')
        return redirect('dashboard:roomtype_list', hotel_pk=hotel.pk)

    return render(request, 'dashboard/rooms/roomtype_form.html', {
        'hotel': hotel, 'room_type': room_type,
        'amenities_all': amenities_all, 'is_edit': is_edit, 'page': 'hotels',
        'selected_amenity_ids': list(room_type.amenities.values_list('id', flat=True)) if room_type.pk else [],
    })


@login_required
@staff_required
def room_by_hotel(request, hotel_pk):
    hotel = get_object_or_404(Hotel, pk=hotel_pk)
    rooms = hotel.rooms.select_related('room_type').order_by('room_number')
    return render(request, 'dashboard/rooms/room_list.html', {
        'hotel': hotel, 'rooms': rooms, 'page': 'hotels',
    })


@login_required
@staff_required
def room_create(request, hotel_pk):
    hotel = get_object_or_404(Hotel, pk=hotel_pk)
    room  = Room(room_type=hotel.room_types.first())
    return _save_room(request, hotel, room)


@login_required
@staff_required
def room_edit(request, hotel_pk, pk):
    hotel = get_object_or_404(Hotel, pk=hotel_pk)
    room  = get_object_or_404(Room, pk=pk, room_type__hotel=hotel)
    return _save_room(request, hotel, room)


@login_required
@staff_required
def room_delete(request, hotel_pk, pk):
    hotel = get_object_or_404(Hotel, pk=hotel_pk)
    room  = get_object_or_404(Room, pk=pk, room_type__hotel=hotel)
    if request.method == 'POST':
        room.delete()
        messages.success(request, f'Đã xóa phòng {room.room_number}.')
    return redirect('dashboard:room_by_hotel', hotel_pk=hotel_pk)


@login_required
@staff_required
def room_bulk_create(request, hotel_pk):
    hotel = get_object_or_404(Hotel, pk=hotel_pk)
    room_types = hotel.room_types.filter(is_active=True)

    if request.method == 'POST':
        room_type_id = request.POST.get('room_type')
        prefix       = request.POST.get('prefix', '').strip()
        start        = int(request.POST.get('start_number', 1))
        count        = int(request.POST.get('count', 1))
        floor        = request.POST.get('floor', '').strip()
        status       = request.POST.get('status', 'available')

        room_type = get_object_or_404(RoomType, pk=room_type_id, hotel=hotel)
        created = 0
        for i in range(count):
            num = str(start + i).zfill(3)
            room_number = f'{prefix}{num}'
            if not Room.objects.filter(room_type__hotel=hotel, room_number=room_number).exists():
                Room.objects.create(
                    room_type=room_type,
                    room_number=room_number,
                    floor=floor,
                    status=status,
                )
                created += 1
        messages.success(request, f'Đã tạo {created} phòng mới.')
        return redirect('dashboard:room_by_hotel', hotel_pk=hotel_pk)

    return render(request, 'dashboard/rooms/room_bulk.html', {
        'hotel': hotel, 'room_types': room_types, 'page': 'hotels',
    })


def _save_room(request, hotel, room):
    room_types = hotel.room_types.filter(is_active=True)
    is_edit    = bool(room.pk)

    if request.method == 'POST':
        rt_id = request.POST.get('room_type')
        room.room_type = get_object_or_404(RoomType, pk=rt_id, hotel=hotel)
        room.room_number = request.POST.get('room_number', '').strip()
        room.floor       = request.POST.get('floor', '').strip()
        room.status      = request.POST.get('status', 'available')
        room.note        = request.POST.get('note', '').strip()

        try:
            room.save()
        except Exception as e:
            messages.error(request, f'Lỗi lưu phòng: {e}')
            return render(request, 'dashboard/rooms/room_form.html', {
                'hotel': hotel, 'room': room, 'room_types': room_types,
                'is_edit': is_edit, 'page': 'hotels',
            })

        action = 'Cập nhật' if is_edit else 'Thêm mới'
        messages.success(request, f'{action} phòng {room.room_number} thành công.')
        return redirect('dashboard:room_by_hotel', hotel_pk=hotel.pk)

    return render(request, 'dashboard/rooms/room_form.html', {
        'hotel': hotel, 'room': room, 'room_types': room_types,
        'is_edit': is_edit, 'page': 'hotels',
    })


# ═══════════════════════════════════════════════════════════
# API STATS
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def api_stats(request):
    today = date.today()
    return JsonResponse({
        'bookings_today': Booking.objects.filter(created_at__date=today).count(),
        'revenue_today': float(
            Payment.objects.filter(status='paid', paid_at__date=today)
            .aggregate(total=Sum('amount'))['total'] or 0
        ),
        'pending': Booking.objects.filter(status='pending').count(),
        'checked_in': Booking.objects.filter(status='checked_in').count(),
    })


# ═══════════════════════════════════════════════════════════
# STAFF — Quản lý nhân viên
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def staff_list(request):
    q = request.GET.get('q', '')
    hotel = request.GET.get('hotel', '')
    staff_qs = StaffProfile.objects.select_related('user', 'hotel').order_by('-created_at')

    if q:
        staff_qs = staff_qs.filter(
            Q(user__username__icontains=q) |
            Q(user__first_name__icontains=q) |
            Q(user__email__icontains=q)
        )
    if hotel:
        staff_qs = staff_qs.filter(hotel_id=hotel)

    hotels_qs = Hotel.objects.filter(is_active=True).order_by('name')
    return render(request, 'dashboard/staff/list.html', {
        'staff_list': staff_qs, 'q': q, 'hotel': hotel,
        'hotels': hotels_qs, 'page': 'staff',
    })


@login_required
@staff_required
def staff_create(request):
    hotels_qs = Hotel.objects.filter(is_active=True).order_by('name')
    if request.method == 'POST':
        return _save_staff(request, None)
    return render(request, 'dashboard/staff/form.html', {
        'hotels': hotels_qs, 'is_edit': False, 'page': 'staff',
    })


@login_required
@staff_required
def staff_edit(request, pk):
    staff = get_object_or_404(StaffProfile, pk=pk)
    hotels_qs = Hotel.objects.filter(is_active=True).order_by('name')
    if request.method == 'POST':
        return _save_staff(request, staff)
    return render(request, 'dashboard/staff/form.html', {
        'staff': staff, 'hotels': hotels_qs, 'is_edit': True, 'page': 'staff',
    })


def _save_staff(request, staff_profile):
    hotels_qs = Hotel.objects.filter(is_active=True).order_by('name')
    is_edit   = staff_profile is not None

    username   = request.POST.get('username', '').strip()
    email      = request.POST.get('email', '').strip()
    first_name = request.POST.get('first_name', '').strip()
    last_name  = request.POST.get('last_name', '').strip()
    hotel_id   = request.POST.get('hotel')
    role       = request.POST.get('role', 'receptionist')
    phone      = request.POST.get('phone', '').strip()
    is_active  = 'is_active' in request.POST

    hotel = get_object_or_404(Hotel, pk=hotel_id)

    if is_edit:
        user = staff_profile.user
        user.email      = email
        user.first_name = first_name
        user.last_name  = last_name
        user.is_active  = is_active
        password = request.POST.get('password', '').strip()
        if password:
            user.set_password(password)
        user.save()
        staff_profile.hotel = hotel
        staff_profile.role  = role
        staff_profile.phone = phone
        staff_profile.save()
        messages.success(request, f'Đã cập nhật nhân viên "{user.username}".')
    else:
        if User.objects.filter(username=username).exists():
            messages.error(request, f'Username "{username}" đã tồn tại.')
            return render(request, 'dashboard/staff/form.html', {
                'hotels': hotels_qs, 'is_edit': False, 'page': 'staff',
            })
        password = request.POST.get('password', '').strip()
        user = User.objects.create_user(
            username=username, email=email,
            first_name=first_name, last_name=last_name,
            password=password, is_active=is_active, is_staff=True,
        )
        # Thêm vào group Employee
        employee_group, _ = Group.objects.get_or_create(name='Employee')
        user.groups.add(employee_group)

        StaffProfile.objects.create(user=user, hotel=hotel, role=role, phone=phone)
        messages.success(request, f'Đã tạo nhân viên "{username}".')

    return redirect('dashboard:staff_list')


@login_required
@staff_required
def staff_delete(request, pk):
    staff = get_object_or_404(StaffProfile, pk=pk)
    if request.method == 'POST':
        user = staff.user
        staff.delete()
        user.delete()
        messages.success(request, 'Đã xóa nhân viên.')
    return redirect('dashboard:staff_list')


@login_required
@staff_required
def staff_toggle(request, pk):
    staff = get_object_or_404(StaffProfile, pk=pk)
    user = staff.user
    if user != request.user:
        user.is_active = not user.is_active
        user.save()
        status = 'kích hoạt' if user.is_active else 'khóa'
        messages.success(request, f'Đã {status} tài khoản nhân viên "{user.username}".')
    return redirect('dashboard:staff_list')


# ═══════════════════════════════════════════════════════════
# SERVICE REQUESTS — Quản lý yêu cầu dịch vụ
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def service_request_list(request):
    from hotels.models import ServiceRequest
    status = request.GET.get('status', '')
    hotel  = request.GET.get('hotel', '')
    reqs   = (
        ServiceRequest.objects
        .select_related('room__room_type__hotel', 'service', 'guest', 'assigned_to')
        .order_by('-created_at')
    )
    if status:
        reqs = reqs.filter(status=status)
    if hotel:
        reqs = reqs.filter(room__room_type__hotel_id=hotel)
    hotels_qs = Hotel.objects.filter(is_active=True).order_by('name')
    return render(request, 'dashboard/service_requests/list.html', {
        'requests': reqs, 'status': status, 'hotel': hotel,
        'hotels': hotels_qs, 'page': 'service_requests',
    })


# ═══════════════════════════════════════════════════════════
# HOMEPAGE — Chỉnh sửa trang chủ (chỉ superuser / staff)
# ═══════════════════════════════════════════════════════════
@login_required
@staff_required
def homepage_editor(request):
    """Chỉnh sửa nội dung hero + section trang chủ."""
    config = HomepageConfig.get()

    if request.method == 'POST':
        config.hero_tagline       = request.POST.get('hero_tagline', '').strip()
        config.hero_title         = request.POST.get('hero_title', '').strip()
        config.hero_subtitle      = request.POST.get('hero_subtitle', '').strip()
        config.hero_image_url     = request.POST.get('hero_image_url', '').strip()
        config.hero_cta_text      = request.POST.get('hero_cta_text', '').strip()
        config.showcase_title     = request.POST.get('showcase_title', '').strip()
        config.showcase_subtitle  = request.POST.get('showcase_subtitle', '').strip()
        config.promo_title        = request.POST.get('promo_title', '').strip()
        config.promo_body         = request.POST.get('promo_body', '').strip()
        config.site_announcement  = request.POST.get('site_announcement', '').strip()
        config.save()
        messages.success(request, '✅ Đã cập nhật trang chủ thành công!')
        return redirect('dashboard:homepage_editor')

    return render(request, 'dashboard/homepage/editor.html', {
        'config': config,
        'page':   'homepage',
    })


# ═══════════════════════════════════════════════════════════
# HOTEL SERVICES — Quản lý dịch vụ khách sạn
# ═══════════════════════════════════════════════════════════

@login_required
@staff_required
def hotel_service_list(request, hotel_pk):
    """Danh sách dịch vụ của một khách sạn."""
    hotel    = get_object_or_404(Hotel, pk=hotel_pk)
    services = hotel.services.order_by('category', 'order', 'name')
    return render(request, 'dashboard/services/list.html', {
        'hotel': hotel, 'services': services, 'page': 'hotels',
    })


@login_required
@staff_required
def hotel_service_create(request, hotel_pk):
    hotel = get_object_or_404(Hotel, pk=hotel_pk)
    if request.method == 'POST':
        return _save_service(request, hotel, None)
    return render(request, 'dashboard/services/form.html', {
        'hotel': hotel, 'page': 'hotels', 'action': 'Thêm dịch vụ',
        'categories': HotelService.CATEGORY_CHOICES,
    })


@login_required
@staff_required
def hotel_service_edit(request, hotel_pk, pk):
    hotel   = get_object_or_404(Hotel, pk=hotel_pk)
    service = get_object_or_404(HotelService, pk=pk, hotel=hotel)
    if request.method == 'POST':
        return _save_service(request, hotel, service)
    return render(request, 'dashboard/services/form.html', {
        'hotel': hotel, 'service': service, 'page': 'hotels',
        'action': 'Sửa dịch vụ',
        'categories': HotelService.CATEGORY_CHOICES,
    })


@login_required
@staff_required
def hotel_service_delete(request, hotel_pk, pk):
    hotel   = get_object_or_404(Hotel, pk=hotel_pk)
    service = get_object_or_404(HotelService, pk=pk, hotel=hotel)
    if request.method == 'POST':
        name = service.name
        service.delete()
        messages.success(request, f'Đã xóa dịch vụ "{name}".')
    return redirect('dashboard:hotel_service_list', hotel_pk=hotel_pk)


@login_required
@staff_required
def hotel_service_toggle(request, hotel_pk, pk):
    hotel   = get_object_or_404(Hotel, pk=hotel_pk)
    service = get_object_or_404(HotelService, pk=pk, hotel=hotel)
    if request.method == 'POST':
        service.is_available = not service.is_available
        service.save()
        state = 'bật' if service.is_available else 'tắt'
        messages.success(request, f'Đã {state} dịch vụ "{service.name}".')
    return redirect('dashboard:hotel_service_list', hotel_pk=hotel_pk)


def _save_service(request, hotel, service):
    data = request.POST
    name = data.get('name', '').strip()
    if not name:
        messages.error(request, 'Tên dịch vụ không được để trống.')
        return redirect('dashboard:hotel_service_list', hotel_pk=hotel.pk)

    price_raw = data.get('price', '').strip()
    try:
        price = float(price_raw) if price_raw else None
    except ValueError:
        price = None

    eta_raw = data.get('eta_minutes', '').strip()
    try:
        eta = int(eta_raw) if eta_raw else None
    except ValueError:
        eta = None

    fields = {
        'hotel':         hotel,
        'name':          name,
        'category':      data.get('category', 'other'),
        'icon':          data.get('icon', '✨').strip() or '✨',
        'description':   data.get('description', '').strip(),
        'price':         price,
        'eta_minutes':   eta,
        'order':         int(data.get('order', 0) or 0),
        'is_available':  'is_available' in data,
        'requires_note': 'requires_note' in data,
    }

    if service:
        for k, v in fields.items():
            setattr(service, k, v)
        service.save()
        messages.success(request, f'✅ Đã cập nhật dịch vụ "{name}".')
    else:
        HotelService.objects.create(**fields)
        messages.success(request, f'✅ Đã thêm dịch vụ "{name}".')

    return redirect('dashboard:hotel_service_list', hotel_pk=hotel.pk)